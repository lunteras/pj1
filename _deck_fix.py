import os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

DECK = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "artifacts", "보장only_추가가입_예측_1to3단계.pptx"))
shutil.copy2(DECK, DECK.replace(".pptx", "_backup12.pptx"))
prs = Presentation(DECK)

# --- 1) S9 축소: E·F 행 제거 + 제목 변경 ---
s9 = list(prs.slides)[8]
for sh in s9.shapes:
    if sh.has_table:
        t = sh.table
        removed = 0
        for row in list(t.rows)[1:]:
            first = row.cells[0].text_frame.text.strip()
            if (first.startswith("E") and "특약" in first) or (first.startswith("F") and "납입" in first):
                t._tbl.remove(row._tr); removed += 1
        print("S9 행 제거:", removed)
for sh in s9.shapes:
    if sh.has_text_frame and "새 가설" in sh.text_frame.text:
        tf = sh.text_frame
        p = tf.paragraphs[0]
        if p.runs:
            p.runs[0].text = "새 가설 A~D — 지지 2 · 반박 2 (E·F는 다음 장에서 강도 점검)"
            for r in p.runs[1:]: r.text = ""
        else:
            tf.text = "새 가설 A~D — 지지 2 · 반박 2 (E·F는 다음 장에서 강도 점검)"
        print("S9 제목 변경 완료")
        break
